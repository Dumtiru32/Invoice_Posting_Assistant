#!/usr/bin/env python3
"""Excel report inspection and ordered document extraction for PDF merging."""
from __future__ import annotations

from collections import Counter
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

STANDARD_ACCOUNTS = {"609007", "607598"}
SPECIAL_ACCOUNTS = {"601101", "607598"}
WASTE_ACCOUNT = "607599"
SEGMENT4_COLUMN = 3
EXTERNAL_DOCUMENT_COLUMN = 6
DEFAULT_DESCRIPTION_COLUMN = 8
OVERVIEW_SHEET = "Overview"


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_document_number(value: Any) -> str:
    text = clean(value)
    if not text or text.casefold() == "(blank)":
        return ""
    return text.replace("/", "-")


def canonical_lookup(value: Any) -> str:
    """Case-insensitive filename key while preserving letters and digits."""
    return re.sub(r"[^a-z0-9]+", "", normalize_document_number(value).casefold())


def account_code(text: Any) -> str:
    """Return the 6-digit account code from labels like '607599', '607599 Total'
    or '607599 - Waste'. Falls back to the trimmed text if no code is found."""
    value = clean(text)
    match = re.match(r"\s*(\d{6})(?!\d)", value)
    return match.group(1) if match else value


def is_total_label(text: Any) -> bool:
    return clean(text).casefold().endswith(" total")


def find_header_row(ws) -> int:
    for row in range(1, ws.max_row + 1):
        segment_header = clean(ws.cell(row=row, column=SEGMENT4_COLUMN).value)
        document_header = clean(ws.cell(row=row, column=EXTERNAL_DOCUMENT_COLUMN).value)
        if (
            re.sub(r"\s+", "", segment_header).casefold() == "segment4"
            and re.sub(r"\s+", "", document_header).casefold()
            in {"externaldocumentnr.", "externaldocumentnr"}
        ):
            return row
    raise ValueError("Could not find the header row. Expected Segment4 in column C and External Document Nr. in column F.")


def find_column_by_header(ws, header_row: int, names: set[str]) -> int | None:
    for column in range(1, ws.max_column + 1):
        header = re.sub(r"\s+", "", clean(ws.cell(header_row, column).value)).casefold()
        if header in names:
            return column
    return None


def find_remarks_column(ws, header_row: int) -> int | None:
    return find_column_by_header(ws, header_row, {"remark", "remarks", "comment", "comments"})


def find_description_column(ws, header_row: int) -> int:
    return find_column_by_header(ws, header_row, {"description"}) or DEFAULT_DESCRIPTION_COLUMN


def detect_special_customer(ws, workbook_path: Path) -> bool:
    search_parts = [workbook_path.stem]
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True):
        search_parts.extend(clean(value) for value in row if value is not None)
    haystack = " ".join(search_parts).casefold()
    return "delonghi" in haystack or "kitchena" in haystack


def inspect_workbook(workbook_path: str | Path) -> dict:
    """Read the Overview pivot-style report, carrying parent labels downward.

    Excel pivot reports commonly display Segment4 and External Document Nr. only
    on the first row of a group. Blank cells underneath mean "same value as
    above", not "no value". Total rows reset the appropriate carried values.

    Account 607599 (waste) is handled separately: it never produces entries
    that need a matching PDF by document number. If the account is present,
    the caller must ask the user to upload ALL waste PDFs (count unknown).
    """
    workbook_path = Path(workbook_path)
    if workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("The report must be an .xlsx or .xlsm workbook.")
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"The Excel report could not be opened: {exc}") from exc
    try:
        if OVERVIEW_SHEET not in workbook.sheetnames:
            raise ValueError('The workbook does not contain the required "Overview" sheet.')
        ws = workbook[OVERVIEW_SHEET]
        header_row = find_header_row(ws)
        remarks_column = find_remarks_column(ws, header_row)
        description_column = find_description_column(ws, header_row)
        special_customer = detect_special_customer(ws, workbook_path)
        allowed_accounts = SPECIAL_ACCOUNTS if special_customer else STANDARD_ACCOUNTS

        entries = []
        waste_documents = []      # waste rows that DO carry an external document nr.
        waste_descriptions = []   # description hints for every waste row
        waste_account_present = False
        empty_document_rows = []
        current_account = ""
        current_document = ""

        for row_number in range(header_row + 1, ws.max_row + 1):
            raw_account = clean(ws.cell(row=row_number, column=SEGMENT4_COLUMN).value)
            raw_document = clean(ws.cell(row=row_number, column=EXTERNAL_DOCUMENT_COLUMN).value)

            # Detect waste before any subtotal row is skipped. A collapsed
            # PivotTable may expose the account only as "607599 Total".
            if raw_account and account_code(raw_account) == WASTE_ACCOUNT:
                waste_account_present = True

            # Pivot subtotal/total rows terminate inherited context.
            if is_total_label(raw_account):
                current_account = ""
                current_document = ""
                continue
            if is_total_label(raw_document):
                current_document = ""
                continue

            # Segment4 is a parent pivot label. Carry it down through child rows.
            if raw_account:
                code = account_code(raw_account)
                if code != current_account:
                    current_document = ""  # a new account starts a new document grouping
                current_account = code
            account = current_account

            # ---- Waste (607599): never goes to entries / empty_document_rows ----
            if account == WASTE_ACCOUNT:
                waste_account_present = True
                description = clean(ws.cell(row=row_number, column=description_column).value)
                if description and description.casefold() != "(blank)":
                    waste_descriptions.append({"row": row_number, "description": description})
                document = normalize_document_number(raw_document)
                if document:
                    remarks = clean(ws.cell(row=row_number, column=remarks_column).value) if remarks_column else ""
                    waste_documents.append({
                        "sequence": len(waste_documents) + 1,
                        "row": row_number,
                        "documentNumber": document,
                        "remarks": remarks,
                    })
                continue

            if account not in allowed_accounts:
                continue

            # External Document Nr. is also a pivot parent label. A requirement is
            # created only on rows where a new document number is explicitly
            # printed, so child description rows do not create duplicates.
            if not raw_document:
                continue
            current_document = normalize_document_number(raw_document)
            document = current_document
            if not document:
                empty_document_rows.append(row_number)
                continue

            remarks = clean(ws.cell(row=row_number, column=remarks_column).value) if remarks_column else ""
            entries.append({
                "sequence": len(entries) + 1,
                "row": row_number,
                "account": account,
                "documentNumber": document,
                "rawDocumentNumber": raw_document,
                "remarks": remarks,
                "lookupKey": canonical_lookup(document),
                "requiresExplicitUpload": False,
                "isWasteDocument": False,
                "documentType": "Invoice document",
            })

        if not entries and not waste_account_present:
            raise ValueError("No qualifying rows were found in the Overview sheet for the configured Segment4 accounts.")

        # A document may occur on several Pivot rows. One physical PDF is required
        # and merged only once, at the position of its first occurrence.
        counts = Counter(entry["lookupKey"] for entry in entries)
        duplicate_keys = {key for key, count in counts.items() if key and count > 1}
        duplicate_documents = sorted({entry["documentNumber"] for entry in entries if entry["lookupKey"] in duplicate_keys})
        duplicate_document_rows = {
            key: [entry["row"] for entry in entries if entry["lookupKey"] == key]
            for key in duplicate_keys
        }
        unique_entries = []
        seen_document_keys = set()
        for entry in entries:
            key = entry["lookupKey"]
            if key in seen_document_keys:
                continue
            seen_document_keys.add(key)
            entry["sequence"] = len(unique_entries) + 1
            if key in duplicate_document_rows:
                entry["sourceRows"] = duplicate_document_rows[key]
                entry["occurrenceCount"] = len(duplicate_document_rows[key])
            unique_entries.append(entry)
        entries = unique_entries

        warnings = []
        if duplicate_documents:
            warnings.append(
                "Repeated document references were consolidated. Upload one PDF per unique document; "
                "each PDF will be merged once at its first Pivot-table occurrence."
            )
        #!if empty_document_rows:
            #!warnings.append("Some qualifying rows have a blank External Document Nr.")
        if waste_account_present:
            warnings.append(
                "Account 607599 (waste) is present. Upload ALL waste-related PDFs; "
                "the number of documents cannot be determined from the report."
            )

        return {
            "workbookName": workbook_path.name,
            "outputName": workbook_path.with_suffix(".pdf").name,
            "sheet": OVERVIEW_SHEET,
            "headerRow": header_row,
            "remarksColumn": remarks_column,
            "descriptionColumn": description_column,
            "specialCustomerRule": special_customer,
            "allowedAccounts": sorted(allowed_accounts | {WASTE_ACCOUNT}),
            "entries": entries,
            "count": len(entries),
            "batchSize": 20,
            "batchCount": (len(entries) + 19) // 20,
            "duplicateDocuments": duplicate_documents,
            "emptyDocumentRows": empty_document_rows,
            # Waste: count is unknown by design; the UI must accept multiple files.
            "wasteDocumentsRequired": waste_account_present,
            "wasteDocumentCount": None,
            "wasteDocumentCountKnown": False,
            "wasteDocuments": waste_documents,          # only hints, may be empty
            "wasteDescriptions": waste_descriptions,    # e.g. "Renewi 8692893/8692894 2026/08"
            "explicit607599Documents": waste_documents,  # kept for backward compatibility
            "warnings": warnings,
        }
    finally:
        workbook.close()
